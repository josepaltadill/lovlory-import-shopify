import csv
import html
import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "datos/03-archivos-importacion-shopify/orgie-tenga-svakom-shunga/productos de Tenga, Shunga, Svakom i Orgie.xlsx"
OLD_MVP = ROOT / "datos/03-archivos-importacion-shopify/shopify_import_lovlory_mvp_productos_con_pesos.csv"
SAINTSUAL = ROOT / "datos/03-archivos-importacion-shopify/saintsual/saintsual-import-silexd-mapale-shunga.csv"
ORGIE_EXPORTS = [
    ROOT / ".codex-tmp/shopify-current-orgie-page1-20260915/products_export.csv",
    ROOT / ".codex-tmp/shopify-current-orgie-page2-20260915/products_export.csv",
]


def clean(value):
    return "" if value is None else str(value).strip()


def norm(value):
    value = html.unescape(clean(value)).casefold()
    value = unicodedata.normalize("NFKD", value)
    value = "".join(char for char in value if not unicodedata.combining(char))
    value = re.sub(r"\b(orgie|tenga|svakom|shunga)\b", " ", value)
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return " ".join(value.split())


def read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook.active
raw_headers = [clean(cell.value) for cell in next(sheet.iter_rows())]
source_rows = []
for values in sheet.iter_rows(values_only=True):
    row = {raw_headers[index]: clean(value) for index, value in enumerate(values)}
    row["Marca"] = row["Marca"].upper()
    row["_images"] = clean(values[7])
    row["_category"] = clean(values[6])
    row["_pvp"] = clean(values[14])
    row["_weight"] = clean(values[12])
    row["_stock"] = clean(values[18])
    source_rows.append(row)

existing = []
for path in ORGIE_EXPORTS:
    existing.extend(read_csv(path))
existing.extend(read_csv(OLD_MVP))
existing.extend(read_csv(SAINTSUAL))

existing_skus = {clean(row.get("Variant SKU") or row.get("SKU")) for row in existing}
existing_eans = {
    clean(row.get("Variant Barcode") or row.get("Barcode"))
    for row in existing
    if clean(row.get("Variant Barcode") or row.get("Barcode"))
}
existing_titles = {norm(row.get("Title")) for row in existing if clean(row.get("Title"))}

tenga_duplicate_skus = {
    "130313",  # 3D Polygon
    "130530",  # Flex White
    "130473",  # Flex Fizzy Green
    "130480",  # Flip 0 Electronic Vibration
    "130569",  # Flip Zero Red + warmer
    "130548",  # Geo Glacier
    "130567",  # Aero Silver
    "130568",  # Aero Cobalt
    "130448",  # Pocket Click Ball
    "130451",  # Pocket Crystal Mist
    "130452",  # Pocket Spark Beads
    "130585",  # UNI Emerald
    "130586",  # UNI Diamond
    "130720",  # Lotion Light
}

results = []
for row in source_rows:
    if row["Marca"] not in {"ORGIE", "SHUNGA", "SVAKOM", "TENGA"}:
        continue
    sku = row["SKU"]
    ean = row["EAN"]
    title = row["Nombre"]
    reason = ""
    if sku and sku in existing_skus:
        reason = "sku"
    elif ean and ean in existing_eans:
        reason = "ean"
    elif norm(title) and norm(title) in existing_titles:
        reason = "title"
    elif row["Marca"] == "TENGA" and sku in tenga_duplicate_skus:
        reason = "tenga_model"
    results.append({**row, "duplicate_reason": reason, "is_new": not bool(reason)})

summary = {}
for brand in ("ORGIE", "SHUNGA", "SVAKOM", "TENGA"):
    brand_rows = [row for row in results if row["Marca"] == brand]
    summary[brand] = {
        "source": len(brand_rows),
        "new": sum(row["is_new"] for row in brand_rows),
        "duplicates": sum(not row["is_new"] for row in brand_rows),
        "duplicate_reasons": {
            reason: sum(row["duplicate_reason"] == reason for row in brand_rows)
            for reason in ("sku", "ean", "title", "tenga_model")
        },
    }

print(json.dumps(summary, indent=2, ensure_ascii=False))
(ROOT / ".codex-tmp/duplicate-summary.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
)
new_rows = [row for row in results if row["is_new"]]
importable_rows = [
    row
    for row in new_rows
    if norm(row["_category"]) != "material pos"
    and not re.search(r"\b(tester|display|expositor)\b", norm(row["Nombre"]))
]
new_urls = []
for row in new_rows:
    new_urls.extend(
        url.strip() for url in row["_images"].split(",") if url.strip()
    )
print(
    json.dumps(
        {
            "new_total": len(new_rows),
            "new_missing_images": sum(not row["_images"] for row in new_rows),
            "new_missing_recommended_price": sum(
                not row["_pvp"] for row in new_rows
            ),
            "new_missing_product_weight": sum(
                not row["_weight"] for row in new_rows
            ),
            "new_zero_stock": sum(
                not row["_stock"] or float(row["_stock"].replace(",", ".")) <= 0
                for row in new_rows
            ),
            "new_image_urls": len(new_urls),
            "new_unique_image_urls": len(set(new_urls)),
            "new_http_image_urls": sum(url.startswith("http://") for url in new_urls),
            "new_https_image_urls": sum(url.startswith("https://") for url in new_urls),
        },
        indent=2,
        ensure_ascii=False,
    )
)
print("\nNEW CATEGORIES")
for brand in ("ORGIE", "SHUNGA", "SVAKOM", "TENGA"):
    print(brand, Counter(row["_category"] for row in new_rows if row["Marca"] == brand))
(ROOT / ".codex-tmp/new-products.json").write_text(
    json.dumps(
        [
            {
                "sku": row["SKU"],
                "title": row["Nombre"],
                "description": row[raw_headers[2]],
                "ean": row["EAN"],
                "brand": row["Marca"],
                "cost": row[raw_headers[5]],
                "category": row["_category"],
                "images": row["_images"],
                "features": row[raw_headers[8]],
                "battery": row[raw_headers[9]],
                "material": row[raw_headers[10]],
                "color": row[raw_headers[11]],
                "product_weight": row["_weight"],
                "package_weight": row[raw_headers[13]],
                "recommended_price": row["_pvp"],
                "dimensions": row[raw_headers[15]],
                "package_dimensions": row[raw_headers[16]],
                "ingredients": row[raw_headers[17]],
                "stock": row["_stock"],
                "registration_date": row[raw_headers[19]],
            }
            for row in importable_rows
        ],
        ensure_ascii=False,
        indent=2,
    ),
    encoding="utf-8",
)
print(f"\nIMPORTABLE EXCLUDING POS/TESTERS/DISPLAYS: {len(importable_rows)}")
(ROOT / ".codex-tmp/new-supplier-image-urls.txt").write_text(
    "\n".join(dict.fromkeys(new_urls)), encoding="utf-8"
)
for brand in ("ORGIE", "SVAKOM"):
    print(f"\nNEW {brand}")
    for row in results:
        if row["Marca"] == brand and row["is_new"]:
            print(row["SKU"], row["EAN"], row["Nombre"])
