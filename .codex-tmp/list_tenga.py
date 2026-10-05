from pathlib import Path

import openpyxl


path = Path(
    "datos/03-archivos-importacion-shopify/orgie-tenga-svakom-shunga/"
    "productos de Tenga, Shunga, Svakom i Orgie.xlsx"
)
workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
sheet = workbook.active
headers = [cell.value for cell in next(sheet.iter_rows())]
column = {name: index for index, name in enumerate(headers)}

selected = {
    "130313", "130315", "130742", "130530", "130473", "130480",
    "130569", "130548", "130567", "130568", "130448", "130451",
    "130452", "130585", "130586", "130720", "130724",
}
for row in sheet.iter_rows(values_only=True):
    if str(row[column["Marca"]]).strip().casefold() != "tenga":
        continue
    if str(row[column["SKU"]]) in selected:
        print(
            f'--- {row[column["SKU"]]} | {row[column["EAN"]]} | '
            f'{row[1]} | {row[5]} | {row[14]} | {row[15]}\n'
            f'{row[2]}\n'
        )
